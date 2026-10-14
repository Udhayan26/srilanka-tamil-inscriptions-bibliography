"""Build the categorised Excel workbook and CSV from data/raw/*.tsv.

Usage:  python scripts/build_excel.py
Output: output/SriLanka_Tamil_Inscriptions_Bibliography.xlsx
        output/bibliography.csv
"""
import csv
import difflib
import glob
import os
import re
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RAW = os.path.join(ROOT, "data", "raw")
OUT = os.path.join(ROOT, "output")
WINDOW_START, WINDOW_END = 1996, 2026

CATEGORIES = [
    ("Books", "Books"),
    ("Research Papers", "Research Papers"),
    ("Newspaper Articles", "Newspaper Articles"),
    ("Magazines/Journals", "Magazines & Journals"),
    ("Web Articles", "Web Articles"),
]

# Theme keywords (Tamil + English) used to tag every record.
THEMES = {
    "Inscriptions / கல்வெட்டு": ["inscription", "கல்வெட்டு", "epigraph", "சாசன", "copper plate", "seal", "slab", "pillar"],
    "Brahmi / பிராமி": ["brahmi", "பிராமி", "potsherd", "graffiti"],
    "Archaeology / தொல்லியல்": ["archaeolog", "தொல்லியல்", "தொல்பொருள்", "excavation", "அகழ்வ", "megalith", "பெருங்கற்", "iron age", "settlement", "monument", "heritage", "ruins"],
    "Hero stones / நடுகல்": ["hero stone", "நடுகல்", "nadukal", "memorial stone"],
    "Rock art / பாறை ஓவியம்": ["rock art", "rock painting", "பாறை ஓவிய", "painting", "ஓவிய"],
    "Coins / நாணயம்": ["coin", "நாணய", "numismat"],
}


def norm_title(t):
    t = t or ""
    stripped = re.sub(r"\(.*?\)", "", t)
    if re.sub(r"[\W_]+", "", stripped):  # keep parenthesised text when it is the whole title
        t = stripped
    return re.sub(r"[\W_]+", " ", t.lower()).strip()


def norm_url(u):
    u = (u or "").strip().replace("http://", "https://").replace("://www.", "://")
    return u.rstrip("/").lower()


def filled(rec):
    return sum(1 for v in rec.values() if v)


def load_rows():
    rows = []
    for path in sorted(glob.glob(os.path.join(RAW, "*.tsv"))):
        with open(path, encoding="utf-8") as f:
            for r in csv.DictReader(f, delimiter="\t"):
                r = {k: (v or "").strip() for k, v in r.items() if k}
                if r.get("title"):
                    r["source_file"] = os.path.basename(path)
                    rows.append(r)
    return rows


def dedupe(rows):
    # URLs cited by many rows are profile/listing pages; only article-like URLs identify a record
    url_uses = {}
    for r in rows:
        url_uses[norm_url(r["url"])] = url_uses.get(norm_url(r["url"]), 0) + 1
    kept = []
    for r in rows:
        dup = None
        a = norm_title(r["title"])
        for k in kept:
            b = norm_title(k["title"])
            if not a or not b:
                continue
            same_url = r["url"] and norm_url(r["url"]) == norm_url(k["url"])
            sim = difflib.SequenceMatcher(None, a, b).ratio()
            auto_pair = r["source_file"].startswith("auto_") or k["source_file"].startswith("auto_")
            article_url = same_url and url_uses.get(norm_url(r["url"]), 0) <= 2 and r["category"] == k["category"]
            if (same_url and sim > 0.55) or article_url or (sim > 0.9 and (r["category"] == k["category"] or auto_pair)):
                dup = k
                break
        if dup is None:
            kept.append(r)
        else:
            # hand-collected rows win over auto-collected ones; otherwise keep the more complete record
            r_auto, d_auto = r["source_file"].startswith("auto_"), dup["source_file"].startswith("auto_")
            if r_auto != d_auto:
                best, other = (dup, r) if r_auto else (r, dup)
            else:
                best, other = (r, dup) if filled(r) > filled(dup) else (dup, r)
            skip = {"notes", "source_file", "category"} if other["source_file"].startswith("auto_") else {"source_file"}
            for key, val in other.items():
                if key not in skip and not best.get(key) and val:
                    best[key] = val
            kept[kept.index(dup)] = best
    return kept


def themes_for(rec):
    hay = " ".join([rec["title"], rec["keywords"], rec["type"]]).lower()
    return "; ".join(name for name, kws in THEMES.items() if any(kw.lower() in hay for kw in kws))


def in_window(rec):
    y = rec["year"][:4]
    if not y.isdigit():
        return "Unknown"
    return "Yes" if WINDOW_START <= int(y) <= WINDOW_END else "No (earlier reference)"


COLUMNS = [
    ("ID", 9), ("Year", 7), ("Date", 11), ("Title", 60), ("Author / Editor", 28),
    ("Publisher / Source", 34), ("Type", 18), ("Language", 10), ("Themes", 30),
    ("Keywords", 30), ("Within 1996-2026", 14), ("Notes", 30), ("URL", 45),
]

HEAD_FILL = PatternFill("solid", start_color="1F4E5F")
HEAD_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name="Arial", size=10)
LINK_FONT = Font(name="Arial", size=10, color="0563C1", underline="single")
THIN = Side(style="thin", color="D0D7DE")
BORDER = Border(bottom=THIN)
BAND = PatternFill("solid", start_color="F3F7F9")


def write_table(ws, records, id_prefix, with_category=False):
    cols = list(COLUMNS)
    if with_category:
        cols.insert(1, ("Category", 20))
    wrap = {"Title", "Author / Editor", "Publisher / Source", "Themes", "Keywords", "Notes"}
    for ci, (name, width) in enumerate(cols, 1):
        c = ws.cell(row=1, column=ci, value=name)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(ci)].width = width
    ws.row_dimensions[1].height = 30
    for ri, rec in enumerate(records, 2):
        vals = [
            f"{id_prefix}-{ri - 1:04d}", rec["year"], rec["date"], rec["title"], rec["author"],
            rec["publisher_source"], rec["type"], rec["language"], rec["themes"], rec["keywords"],
            rec["window"], rec["notes"], rec["url"],
        ]
        if with_category:
            vals.insert(1, rec["category"])
        for ci, v in enumerate(vals, 1):
            c = ws.cell(row=ri, column=ci, value=v if v != "" else None)
            c.font = BODY_FONT
            c.alignment = Alignment(vertical="top", wrap_text=cols[ci - 1][0] in wrap)
            c.border = BORDER
            if ri % 2 == 0:
                c.fill = BAND
            if cols[ci - 1][0] == "Year":
                c.number_format = "@"
        if rec["url"]:
            u = ws.cell(row=ri, column=len(cols))
            u.hyperlink, u.font = rec["url"], LINK_FONT
    title_col = [n for n, _ in cols].index("Title") + 1
    ws.freeze_panes = f"{get_column_letter(title_col + 1)}2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(len(records) + 1, 2)}"


def sort_key(rec):
    y = rec["year"][:4]
    return (-(int(y) if y.isdigit() else 0), rec["date"], rec["title"])


def build():
    records = dedupe(load_rows())
    for r in records:
        r["themes"] = themes_for(r)
        r["window"] = in_window(r)
    records.sort(key=sort_key)

    os.makedirs(OUT, exist_ok=True)
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"

    prefixes = {"Books": "BK", "Research Papers": "RP", "Newspaper Articles": "NW", "Magazines/Journals": "MG",
                "Web Articles": "WB"}
    for cat, sheet in CATEGORIES:
        ws = wb.create_sheet(sheet)
        write_table(ws, [r for r in records if r["category"] == cat], prefixes[cat])
    ws_all = wb.create_sheet("All Records")
    all_rows = sorted(records, key=lambda r: ([c for c, _ in CATEGORIES].index(r["category"]), sort_key(r)))
    write_table(ws_all, all_rows, "ALL", with_category=True)

    build_summary(summary, len(all_rows))
    build_sources(wb)

    xlsx = os.path.join(OUT, "SriLanka_Tamil_Inscriptions_Bibliography.xlsx")
    wb.save(xlsx)

    with open(os.path.join(OUT, "bibliography.csv"), "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["category", "year", "date", "title", "author", "publisher_source", "type", "language",
                    "themes", "keywords", "within_1996_2026", "notes", "url"])
        for r in all_rows:
            w.writerow([r["category"], r["year"], r["date"], r["title"], r["author"], r["publisher_source"],
                        r["type"], r["language"], r["themes"], r["keywords"], r["window"], r["notes"], r["url"]])
    print(f"{len(all_rows)} records -> {xlsx}")
    return xlsx


def build_summary(ws, n):
    title_font = Font(name="Arial", bold=True, size=14, color="1F4E5F")
    bold = Font(name="Arial", bold=True, size=10)
    body = Font(name="Arial", size=10)
    ws["A1"] = "Sri Lankan Tamil Inscriptions, Archaeology, Hero Stones & Rock Art: Bibliography"
    ws["A1"].font = title_font
    ws["A2"] = (f"Books, research papers, newspaper and magazine articles, 1996-2026 (plus key earlier references). "
                f"Compiled {date.today().isoformat()}. Counts below are live formulas over the 'All Records' sheet.")
    ws["A2"].font = body
    ws["A3"] = "Keywords: கல்வெட்டு, கல்வெட்டுகள், தமிழ், தொல்லியல், நடுகல், பாறை ஓவியம் / inscription, epigraphy, Tamil, archaeology, hero stone, rock art, Brahmi, coins"
    ws["A3"].font = body

    rng = f"'All Records'!$B$2:$B${n + 1}"
    yr = f"'All Records'!$C$2:$C${n + 1}"
    th = f"'All Records'!$J$2:$J${n + 1}"
    win = f"'All Records'!$L$2:$L${n + 1}"

    ws["A5"], ws["B5"], ws["C5"] = "Category", "Records", "Within 1996-2026"
    for c in ("A5", "B5", "C5"):
        ws[c].font, ws[c].fill = HEAD_FONT, HEAD_FILL
    for i, (cat, _) in enumerate(CATEGORIES, 6):
        ws.cell(row=i, column=1, value=cat).font = body
        ws.cell(row=i, column=2, value=f'=COUNTIF({rng},A{i})').font = body
        ws.cell(row=i, column=3, value=f'=COUNTIFS({rng},A{i},{win},"Yes")').font = body
    t = 6 + len(CATEGORIES)
    ws.cell(row=t, column=1, value="Total").font = bold
    ws.cell(row=t, column=2, value=f"=SUM(B6:B{t - 1})").font = bold
    ws.cell(row=t, column=3, value=f"=SUM(C6:C{t - 1})").font = bold

    r0 = t + 2
    ws.cell(row=r0, column=1, value="Theme").font = HEAD_FONT
    ws.cell(row=r0, column=2, value="Records").font = HEAD_FONT
    ws.cell(row=r0, column=1).fill = ws.cell(row=r0, column=2).fill = HEAD_FILL
    for i, name in enumerate(THEMES, r0 + 1):
        ws.cell(row=i, column=1, value=name).font = body
        ws.cell(row=i, column=2, value=f'=COUNTIF({th},"*"&A{i}&"*")').font = body
    ws.cell(row=r0 + len(THEMES) + 1, column=1,
            value="A record can carry several themes, so theme counts add up to more than the total.").font = Font(name="Arial", size=9, italic=True)

    r1 = r0 + len(THEMES) + 3
    hdr = ["Period"] + [c for c, _ in CATEGORIES] + ["All"]
    for ci, h in enumerate(hdr, 1):
        c = ws.cell(row=r1, column=ci, value=h)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
    periods = [("1996-2000", 1996, 2000), ("2001-2005", 2001, 2005), ("2006-2010", 2006, 2010),
               ("2011-2015", 2011, 2015), ("2016-2020", 2016, 2020), ("2021-2026", 2021, 2026)]
    for i, (label, a, b) in enumerate(periods, r1 + 1):
        ws.cell(row=i, column=1, value=label).font = body
        for ci, (cat, _) in enumerate(CATEGORIES, 2):
            ws.cell(row=i, column=ci,
                    value=f'=SUMPRODUCT(({rng}="{cat}")*(IFERROR(VALUE(LEFT({yr},4)),0)>={a})*(IFERROR(VALUE(LEFT({yr},4)),0)<={b}))').font = body
        ws.cell(row=i, column=len(hdr), value=f"=SUM(B{i}:{get_column_letter(len(hdr) - 1)}{i})").font = bold
    note_row = r1 + len(periods) + 1
    ws.cell(row=note_row, column=1,
            value="Records with no year yet are counted in the category totals but not in this period table.").font = Font(name="Arial", size=9, italic=True)

    ws.column_dimensions["A"].width = 30
    for col in "BCDEFG":
        ws.column_dimensions[col].width = 18


def build_sources(wb):
    ws = wb.create_sheet("Sources & Method")
    rows = [
        ("Source", "What it contributed", "URL"),
        ("Noolaham digital library", "Sri Lankan Tamil books, journals, newspapers (digitised)", "https://noolaham.org"),
        ("Ezhuna (எழுநா)", "Serialised Tamil articles on Brahmi inscriptions, Naga inscriptions, Kantarodai archaeology, East coast history", "https://www.ezhunaonline.com"),
        ("University of Jaffna - Dept. of History staff pages", "Publication lists of archaeology/epigraphy staff", "https://www.his.jfn.ac.lk"),
        ("University of Jaffna - Prof. S. Pathmanathan page", "Inscription books and papers", "https://jfn.ac.lk/prof-s-pathmanathan/"),
        ("University repositories (Jaffna, Peradeniya, SEUSL, Kelaniya)", "Theses, journal articles, conference papers", "https://rer.jfn.ac.lk"),
        ("TamilNet / Tamil Guardian", "English news on finds and debates", "https://www.tamilnet.com"),
        ("Virakesari, Thinakkural, Tamil Mirror, Thinakaran, Tamilwin, IBC Tamil", "Tamil newspaper reports and features", "https://www.virakesari.lk"),
        ("Sunday Times, Sunday Observer, Daily News, The Island, Colombo Telegraph", "English newspaper features and opinion", "https://www.sundaytimes.lk"),
        ("Lanka Pradeepa list of Tamil inscriptions", "27-item reference list of Tamil inscription publications", "https://www.lankapradeepa.com/2020/04/tamil-inscriptions-in-sri-lanka_16.html"),
        ("", "", ""),
        ("Method", "Searched Tamil and English keywords (கல்வெட்டு, கல்வெட்டுகள், தமிழ், தொல்லியல், நடுகல், பாறை ஓவியம், inscription, Brahmi, epigraphy, hero stone, rock art, coins), opened catalogue and listing pages, recorded bibliographic details, then removed duplicates.", ""),
        ("Verification", "Rows whose Notes say 'to verify' have a missing or unconfirmed year, author or title; open the URL to confirm.", ""),
        ("Updates", "The GitHub repository runs a daily scraper (news, web pages, Wikipedia, Ezhuna, OpenAlex, Crossref) that appends new items to data/raw/auto_*.tsv and rebuilds this workbook. A one-time backfill searched each year from 1996.", ""),
    ]
    for ri, row in enumerate(rows, 1):
        for ci, v in enumerate(row, 1):
            c = ws.cell(row=ri, column=ci, value=v or None)
            c.font = HEAD_FONT if ri == 1 else Font(name="Arial", size=10, bold=(ci == 1))
            if ri == 1:
                c.fill = HEAD_FILL
            c.alignment = Alignment(vertical="top", wrap_text=True)
            if ci == 3 and v.startswith("http"):
                c.hyperlink, c.font = v, LINK_FONT
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 80
    ws.column_dimensions["C"].width = 50


if __name__ == "__main__":
    build()
